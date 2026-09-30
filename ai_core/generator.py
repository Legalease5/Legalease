import io
import os
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    # Special characters matrum typographic quotes-a remove panna[span_5](start_span)[span_5](end_span)
    text = text.replace('“', '"').replace('”', '"').replace("’", "'")
    return re.sub(r'[^\x00-\x7F]+', '', text)

def format_docx(text: str, doc_type: str, logo_path: str = "Image/Logo.png") -> io.BytesIO:
    doc = Document()
    
    # Header-la Logo insert panna[span_6](start_span)[span_6](end_span)
    if os.path.exists(logo_path):
        section = doc.sections[0]
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.add_run().add_picture(logo_path, width=Inches(1.5))

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.name = 'Times New Roman'

    # Content
    for para in text.split('\n\n'):
        if para.strip():
            p = doc.add_paragraph()
            p_run = p.add_run(para.strip())
            p_run.font.name = 'Times New Roman'
            p_run.font.size = Pt(12)

    # Footer[span_7](start_span)[span_7](end_span)
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.text = "LegalEase Inc. | contact@legalease.com | All Rights Reserved."

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

class PDFGenerator(FPDF):
    def header(self):
        if os.path.exists("Image/Logo.png"):
            self.image("Image/Logo.png", 85, 8, 40)
            self.ln(25)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.cell(0, 10, "LegalEase Inc. | contact@legalease.com | All Rights Reserved.", 0, 0, "C")

def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = PDFGenerator()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, doc_type.upper(), ln=True, align="C")
    pdf.ln(5)

    pdf.set_font("Arial", size=11)
    clean_text = sanitize_text(text)
    for line in clean_text.split("\n"):
        pdf.multi_cell(0, 8, line)

    return pdf.output(dest='S').encode('latin-1')

def format_html_preview(text: str) -> str:
    formatted_html = text.replace("\n", "<br>")
    return f"<div style='background-color: #1e1e1e; color: #ffffff; padding: 20px; border-radius: 8px; font-family: monospace;'>{formatted_html}</div>"